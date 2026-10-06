from io import BytesIO
from pathlib import Path
import re

from docx import Document
from pypdf import PdfReader


MAX_TAILLE_DOCUMENT = 10 * 1024 * 1024
EXTENSIONS_DOCUMENT = {".txt", ".pdf", ".docx"}


def extraire_texte(nom_fichier: str, contenu: bytes) -> str:
    extension = Path(nom_fichier).suffix.lower()

    if extension not in EXTENSIONS_DOCUMENT:
        raise ValueError("Formats acceptés : .txt, .pdf ou .docx.")

    if len(contenu) > MAX_TAILLE_DOCUMENT:
        raise ValueError("Le fichier dépasse la taille maximale de 10 Mo.")

    if extension == ".txt":
        try:
            texte = contenu.decode("utf-8-sig")
        except UnicodeDecodeError as erreur:
            raise ValueError(f"Le fichier {nom_fichier} doit être encodé en UTF-8.") from erreur
    elif extension == ".pdf":
        try:
            document = PdfReader(BytesIO(contenu))
            texte = "\n".join(page.extract_text() or "" for page in document.pages)
        except Exception as erreur:
            raise ValueError(f"Impossible de lire le PDF {nom_fichier}.") from erreur

        if not texte.strip():
            raise ValueError(
                f"Le PDF {nom_fichier} ne contient pas de texte extractible. "
                "Les PDF scannés nécessitent un OCR."
            )
    else:
        try:
            document = Document(BytesIO(contenu))
            paragraphes = [paragraphe.text for paragraphe in document.paragraphs]
            tableaux = [
                "\t".join(cellule.text for cellule in ligne.cells)
                for tableau in document.tables
                for ligne in tableau.rows
            ]
            texte = "\n".join(paragraphes + tableaux)
        except Exception as erreur:
            raise ValueError(f"Impossible de lire le document Word {nom_fichier}.") from erreur

    if not texte.strip():
        raise ValueError(f"Le fichier {nom_fichier} ne contient aucun texte exploitable.")

    return texte.strip()


def extraire_identite(texte: str) -> dict[str, str]:
    nom = re.search(
        r"^\s*Nom(?:\s+de\s+l['’]élève)?\s*[:：]\s*(.*?)\s*$",
        texte,
        re.IGNORECASE | re.MULTILINE,
    )
    prenom = re.search(
        r"^\s*Prénom(?:\s+de\s+l['’]élève)?\s*[:：]\s*(.*?)\s*$",
        texte,
        re.IGNORECASE | re.MULTILINE,
    )

    return {
        "nom": nom.group(1).strip() if nom else "",
        "prenom": prenom.group(1).strip() if prenom else "",
    }